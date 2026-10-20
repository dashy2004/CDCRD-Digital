"""Corrida completa del CLI sobre el ejemplo publicable (datos inventados)."""
import json
import os

from openpyxl import load_workbook

from puente_losas.cli import main

EJEMPLO = os.path.join(os.path.dirname(__file__), "..", "ejemplo")


def test_cli_ejemplo_completo(tmp_path, capsys):
    ret = main(["--proyecto", os.path.join(EJEMPLO, "edificio-ejemplo.json"),
                "--usos", os.path.join(EJEMPLO, "usos-ejemplo.json"),
                "--salida", str(tmp_path), "--nombre", "Edificio Ejemplo PL"])
    assert ret == 0
    assert "6 paños analizados" in capsys.readouterr().out

    wb = load_workbook(str(tmp_path / "informe_losas.xlsx"))
    assert wb["Losas"].max_row == 3 + 6

    plan = json.loads((tmp_path / "plan_etabs.json").read_text(encoding="utf-8"))
    assert plan["meta"]["unidades"] == "Ton, m, C"

    datos = json.loads((tmp_path / "resultados.json").read_text(encoding="utf-8"))
    por_mark = {r["mark"]: r for r in datos["panos"]}
    # La junta LE2|LE3 corta la continuidad: ningún paño queda con 4 vecinos reales.
    assert por_mark["LE3"]["borde"] == "simple"
    # El vuelo quedó sin mampostería y con su set Comun.
    assert por_mark["LE6"]["es_vuelo"] is True
    assert por_mark["LE6"]["cargas_tm2"]["mamposteria_asignada"] == 0.0
    assert por_mark["LE6"]["cargas_tm2"]["viva"] == 0.40
    # La mampostería declarada de LE4 se asignó redondeada hacia arriba.
    assert por_mark["LE4"]["cargas_tm2"]["mamposteria_asignada"] == 0.15
    assert (tmp_path / "mapa_losas.svg").exists()

    # Vía Pieper-Martens: clasificación, momentos y .DL para LosasPlus/Losas.exe.
    pm = datos["pieper_martens"]
    assert pm["losas"]["LE6"]["codigo_dl"] == 72        # vuelo sobre Lx
    assert pm["losas"]["LE1"]["codigo_dl"] in (21, 22, 31, 32, 40, 51, 52)
    assert pm["losas"]["LE5"]["mfx"] > 0
    assert any(a["modo"] == "N" for a in pm["apoyos"])  # el borde del vuelo no balancea
    from puente_losas.dl_io import cargar_dl
    sistema = cargar_dl(str(tmp_path / "sistema_losas.DL"))
    assert len(sistema.losas) == 6
    assert sistema.fc_ton_cm2 == 0.210


def test_cli_version(capsys):
    assert main(["--version"]) == 0
    assert capsys.readouterr().out.strip()


def test_cli_error_limpio(tmp_path, capsys):
    assert main(["--proyecto", str(tmp_path / "no-existe.json")]) == 1
    assert "error:" in capsys.readouterr().err
