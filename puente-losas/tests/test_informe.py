"""Informe Excel y mapa SVG: se escriben y se releen con el contenido esperado."""
from openpyxl import load_workbook

from puente_losas.contrato import PanoLosa, asignar_bordes
from puente_losas.informe_excel import (ENCABEZADOS, escribir_informe,
                                        escribir_mapa_svg)
from puente_losas.motor import analizar_panos


def _resultados():
    panos = [
        PanoLosa(mark="LE1", nivel="N1", z_tope=3.0, lx=5.0, ly=6.0, espesor=0.25,
                 contorno=[[0, 0], [5, 0], [5, 6], [0, 6]]),
        PanoLosa(mark="LE2", nivel="N1", z_tope=3.0, lx=1.8, ly=6.0, espesor=0.25,
                 contorno=[[-1.8, 0], [0, 0], [0, 6], [-1.8, 6]],
                 uso="comun", es_vuelo=True, direccion="1D"),
    ]
    return analizar_panos(asignar_bordes(panos))


def test_informe_excel_relee(tmp_path):
    ruta = tmp_path / "informe.xlsx"
    resultados = _resultados()
    escribir_informe(resultados, str(ruta), proyecto="Edificio Ejemplo")
    wb = load_workbook(str(ruta))
    assert wb.sheetnames == ["Losas", "Supuestos"]
    ws = wb["Losas"]
    assert ws.max_row == 3 + len(resultados)        # título + encabezado + unidades + filas
    assert [c.value for c in ws[2]][:4] == [h for h, _ in ENCABEZADOS][:4]
    marks = [ws.cell(row=4 + i, column=1).value for i in range(len(resultados))]
    assert marks == ["LE1", "LE2"]
    # El vuelo queda marcado y sin mampostería.
    col_vuelo = [h for h, _ in ENCABEZADOS].index("Vuelo") + 1
    col_mamp = [h for h, _ in ENCABEZADOS].index("Mamp. asignada") + 1
    assert ws.cell(row=5, column=col_vuelo).value == "SÍ"
    assert ws.cell(row=5, column=col_mamp).value == 0.0
    # La hoja de supuestos cita el motor y su autoría.
    supuestos = " ".join(str(f[0].value) for f in wb["Supuestos"].iter_rows(min_row=2))
    assert "EstructurasRD" in supuestos and "ACI 318-19" in supuestos


def test_mapa_svg(tmp_path):
    ruta = tmp_path / "mapa.svg"
    escribir_mapa_svg(_resultados(), str(ruta))
    svg = ruta.read_text(encoding="utf-8")
    assert svg.startswith("<svg") and "LE1" in svg and "LE2" in svg
    assert svg.count("<polygon") == 2
