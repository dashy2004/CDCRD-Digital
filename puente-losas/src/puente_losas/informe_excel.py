"""Informe Excel por losa (el entregable del flujo de la Ing. Arlette Ramírez).

Hojas:
- ``Losas``: una fila por paño — luces (crudas y redondeadas ↑0.05 según el
  Anexo A), cargas, momentos, acero por franja, espesor mínimo y chequeos.
- ``Supuestos``: fuentes, unidades y limitaciones declaradas de la corrida.

Unidades declaradas en cada encabezado. Los momentos salen del motor FEM de
EstructurasRD; nada se escribe que no venga de la corrida.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from puente_losas.espesor import redondear_arriba_005_m
from puente_losas.motor import ResultadoPano

ENCABEZADOS = [
    ("Mark", ""), ("Nivel", ""), ("Uso", ""), ("Vuelo", ""),
    ("Lx", "m"), ("Ly", "m"), ("Lx Excel", "m ↑0.05"), ("Ly Excel", "m ↑0.05"),
    ("Dirección", "1D/2D"), ("Borde FEM", ""),
    ("h modelo", "m"), ("h mín norma", "m"), ("h usado", "m"), ("Criterio h mín", ""),
    ("Sección ETABS", ""),
    ("PP", "t/m²"), ("Acabados D", "t/m²"), ("Mamp. asignada", "t/m²"),
    ("Patrón viva", ""), ("Viva", "t/m²"), ("q servicio", "t/m²"), ("q última", "t/m²"),
    ("w servicio", "mm"), ("Límite l/240", "mm"), ("Flecha OK", ""),
    ("Mx vano", "tonf·m/m"), ("My vano", "tonf·m/m"), ("M apoyo", "tonf·m/m"),
    ("As X", "mm²/m"), ("Armadura X", ""), ("As Y", "mm²/m"), ("Armadura Y", ""),
    ("As apoyo", "mm²/m"), ("Armadura apoyo", ""),
    ("Acero OK", ""), ("h modelo suficiente", ""),
    ("Tipo PM", "cód. DL"), ("PM Mfx", "tonf·m/m"), ("PM Mfy", "tonf·m/m"),
    ("PM Msx", "tonf·m/m"), ("PM Msy", "tonf·m/m"),
]

SUPUESTOS = [
    "Motor de cálculo: motor-fea de EstructurasRD (github.com/dashy2004/EstructurasRD, MIT) — "
    "FEM de placa ACM (3 GDL/nodo), carga uniforme lumped, diseño de franjas ACI 318-19 en SI.",
    "Rectángulo equivalente: cada paño se analiza como rectángulo Lx×Ly; huecos y entrantes NO "
    "entran al FEM. Esto es PREDIMENSIONADO: el análisis definitivo es el de ETABS con la "
    "geometría real.",
    "Borde del FEM: 'empotrado' solo si los 4 bordes del paño tienen vecino contiguo (junta de "
    "0.10 m NO cuenta); el resto 'simple'. La condición real mixta queda entre ambos extremos.",
    "Cargas: sets del prompt de la Ing. Arlette Ramírez (14/10/2026) en t/m²; mampostería aparte, "
    "mínimo 0.10 t/m² en paños apoyados, vuelos excluidos de mampostería y del mínimo; redondeo "
    "de mampostería SIEMPRE hacia arriba a 0.05 al asignar.",
    "q última = 1.2D + 1.6L (CDCRD Vol I T02 / ACI 318-19 §5.3.1b). Momentos y acero con q última.",
    "Flecha: w central en servicio (D+L) escalada linealmente de la corrida última; es flecha "
    "ELÁSTICA INMEDIATA con sección bruta (sin fisuración ni fluencia). Límite informativo l/240 "
    "(ACI 318-19 T.24.2.2).",
    "Materiales: C210 (f'c 210 kg/cm²; diseño con 21 MPa por indicación del encargo), "
    "Ec = 15100·√f'c (kg/cm²) = 2 188 198 tonf/m²; acero A615 Gr60 (fy 420 MPa). "
    "Recubrimiento de losa 25 mm.",
    "Espesor mínimo: ACI 318-19 T.7.3.1.1 (1D) o §8.3.1.2(b) con supuesto αfm>2 (2D con vigas); "
    "luces entre ejes como aproximación de luces libres (conservador).",
    "Redondeos del Excel: Lx/Ly SIEMPRE hacia arriba a 0.05 m (Anexo A); nunca hacia abajo.",
    "Columnas PM: método de Pieper-Martens (tablas de F. Perdomo, Losas v5.21) con la q última "
    "del paño y el código de borde deducido de la adyacencia; mismos supuestos de apoyos rígidos "
    "de Losas.exe. La hoja 'Apoyos PM' trae los momentos de borde balanceados (promedio con piso "
    "0.75·max; en vuelos gobierna el vuelo). Validado contra el fixture RESTAURANTE 2 de "
    "EstructurasRD (salida real de Losas.exe).",
]

_CABECERA = Font(bold=True)
_ROJO = PatternFill("solid", start_color="FFC7CE")
_CENTRO = Alignment(horizontal="center")


def _fila(r: ResultadoPano, pm=None) -> list:
    p, c, e = r.pano, r.cargas, r.espesor
    acero_ok = all(f.get("cumple") and not f.get("seccion_insuficiente")
                   for f in (r.franja_x, r.franja_y, r.franja_apoyo))
    if pm is not None and p.mark in pm.momentos:
        m = pm.momentos[p.mark]
        cols_pm = [pm.clasificacion[p.mark].codigo, round(m.mfx, 3), round(m.mfy, 3),
                   round(m.msx, 3), round(m.msy, 3)]
    else:
        cols_pm = ["", "", "", "", ""]
    return [
        p.mark, p.nivel or "", p.uso, "SÍ" if p.es_vuelo else "NO",
        round(p.lx, 3), round(p.ly, 3),
        redondear_arriba_005_m(p.lx), redondear_arriba_005_m(p.ly),
        "1D" if p.trabaja_1d else "2D", p.borde or "",
        round(p.espesor, 3), round(e.h_min_normativo, 4), e.h_usado, e.criterio,
        e.seccion_etabs,
        round(c.peso_propio, 4), c.acabados_dead, c.mamposteria_asignada,
        c.patron_viva, c.viva, round(c.q_servicio, 4), round(c.q_ultima, 4),
        round(r.w_servicio_m * 1000.0, 2), round(r.limite_flecha_m * 1000.0, 1),
        "SÍ" if r.flecha_cumple else "NO",
        round(r.mx_tonf_m, 3), round(r.my_tonf_m, 3), round(r.m_apoyo_tonf_m, 3),
        round(r.franja_x["as_diseno"], 1), r.franja_x["disponer"],
        round(r.franja_y["as_diseno"], 1), r.franja_y["disponer"],
        round(r.franja_apoyo["as_diseno"], 1), r.franja_apoyo["disponer"],
        "SÍ" if acero_ok else "NO",
        "NO" if e.espesor_modelo_insuficiente else "SÍ",
        *cols_pm,
    ]


def escribir_informe(resultados: list[ResultadoPano], ruta: str,
                     proyecto: str = "Proyecto Ejemplo", pm=None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Losas"
    ws.append([f"Informe de losas — {proyecto} — puente-losas (motor EstructurasRD)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([h for h, _ in ENCABEZADOS])
    ws.append([u for _, u in ENCABEZADOS])
    for celda in ws[2] + ws[3]:
        celda.font = _CABECERA
        celda.alignment = _CENTRO
    columnas_no = {ENCABEZADOS.index(("Flecha OK", "")) + 1,
                   ENCABEZADOS.index(("Acero OK", "")) + 1,
                   ENCABEZADOS.index(("h modelo suficiente", "")) + 1}
    for r in resultados:
        ws.append(_fila(r, pm))
        fila = ws.max_row
        for col in columnas_no:
            celda = ws.cell(row=fila, column=col)
            if celda.value == "NO":
                celda.fill = _ROJO
    ws.freeze_panes = "A4"

    if pm is not None and pm.apoyos:
        hoja_pm = wb.create_sheet("Apoyos PM")
        hoja_pm.append(["Momentos balanceados sobre apoyos — Pieper-Martens (ton·m/m)"])
        hoja_pm["A1"].font = Font(bold=True, size=12)
        hoja_pm.append(["Losa I", "Losa J", "Dirección", "Mu I", "Mu J", "Mu I-J", "Balanceo"])
        for celda in hoja_pm[2]:
            celda.font = _CABECERA
        for a in pm.apoyos:
            hoja_pm.append([a.losa_i, a.losa_j, a.direccion, round(a.mu_i, 3),
                            round(a.mu_j, 3), round(a.mu_ij, 3), a.modo])

    hoja = wb.create_sheet("Supuestos")
    hoja.append(["Supuestos y limitaciones de la corrida"])
    hoja["A1"].font = Font(bold=True, size=12)
    for s in SUPUESTOS:
        hoja.append([s])
    hoja.column_dimensions["A"].width = 120
    for fila_ in hoja.iter_rows(min_row=2):
        fila_[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(ruta)


def escribir_mapa_svg(resultados: list[ResultadoPano], ruta: str) -> None:
    """Mapa en planta de los paños, coloreado por estado (verde OK / rojo con falla)."""
    if not resultados:
        raise ValueError("Sin resultados para el mapa.")
    cajas = []
    for r in resultados:
        contorno = r.pano.contorno or [[0, 0], [r.pano.lx, 0], [r.pano.lx, r.pano.ly], [0, r.pano.ly]]
        ok = (r.flecha_cumple and not r.espesor.espesor_modelo_insuficiente and
              all(f.get("cumple") for f in (r.franja_x, r.franja_y, r.franja_apoyo)))
        cajas.append((r.pano.mark, contorno, ok))
    xs = [p[0] for _, c, _ in cajas for p in c]
    ys = [p[1] for _, c, _ in cajas for p in c]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    escala = 900.0 / max(x1 - x0, 1e-6)
    alto = (y1 - y0) * escala + 60

    def px(x: float) -> float:
        return round((x - x0) * escala + 30, 1)

    def py(y: float) -> float:
        return round(alto - 30 - (y - y0) * escala, 1)   # Y arriba, como en planta

    partes = [f'<svg xmlns="http://www.w3.org/2000/svg" width="960" height="{alto:.0f}" '
              f'viewBox="0 0 960 {alto:.0f}" font-family="sans-serif">',
              '<rect width="100%" height="100%" fill="white"/>']
    for mark, contorno, ok in cajas:
        puntos = " ".join(f"{px(p[0])},{py(p[1])}" for p in contorno)
        color = "#8fd19e" if ok else "#f1948a"
        partes.append(f'<polygon points="{puntos}" fill="{color}" stroke="#333" stroke-width="1"/>')
        cx = sum(p[0] for p in contorno) / len(contorno)
        cy = sum(p[1] for p in contorno) / len(contorno)
        partes.append(f'<text x="{px(cx)}" y="{py(cy)}" font-size="14" text-anchor="middle">{mark}</text>')
    partes.append("</svg>")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("\n".join(partes))
