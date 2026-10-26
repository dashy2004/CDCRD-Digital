"""Static, exclusive renderers for an already validated presentation model.

PDF appendices keep their visible content and original pagination. A new bottom
band carries the assembled page number. Only links without an appearance are
removed; visible or unknown annotation types are rejected instead of silently
discarding evidence. Active PDF objects are removed. Word and HTML contain the
editable body and an appendix inventory.
"""

from __future__ import annotations

import re
from html import escape
from io import BytesIO
from math import isfinite, sqrt
from pathlib import Path
from typing import Any

MAX_PDF_PAGES = 2000
MAX_PDF_BYTES = 256 * 1024 * 1024
MAX_TOTAL_PDF_BYTES = 512 * 1024 * 1024
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_STATUS = {"draft": "Borrador", "ready_for_review": "Preparado para revisión"}
_ANNEX_NOTE = (
    "Los anexos completos se incorporan al PDF ensamblado. Este documento "
    "contiene el cuerpo editable y el inventario de anexos."
)


class RenderError(ValueError):
    """An input cannot be represented safely or consistently."""


def _text(value: Any) -> str:
    if not isinstance(value, str):
        raise RenderError("El modelo de presentación debe contener textos explícitos.")
    if _CONTROL.search(value):
        raise RenderError("El texto contiene caracteres de control no admitidos.")
    return value


def _status(view: dict[str, Any]) -> str:
    if view.get("engineering_approval") is not False:
        raise RenderError("El generador no puede declarar aprobación de ingeniería.")
    try:
        return _STATUS[view["document_status"]]
    except (KeyError, TypeError) as exc:
        raise RenderError("Estado documental no admitido.") from exc


def _table_data(table: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    columns = [_text(cell) for cell in table["columns"]]
    if not columns or len(columns) > 8:
        raise RenderError("Una tabla debe tener entre una y ocho columnas.")
    rows = [[_text(cell) for cell in row] for row in table["rows"]]
    if any(len(row) != len(columns) for row in rows):
        raise RenderError("Las filas no coinciden con las columnas de la tabla.")
    return columns, rows


def _widths(table: dict[str, Any], width: float) -> list[float]:
    columns, rows = _table_data(table)
    declared = table.get("widths")
    if declared is not None:
        if len(declared) != len(columns) or any(
            isinstance(x, bool) or not isinstance(x, (int, float))
            or not isfinite(x) or x <= 0 for x in declared
        ):
            raise RenderError("Las proporciones de ancho de tabla no son válidas.")
        weights = list(declared)
    else:
        # Weight descriptions more heavily without letting a long identifier
        # consume the whole page. Every column retains a usable minimum width.
        weights = [sqrt(min(100, max(10, len(columns[i]), *[
            len(row[i]) for row in rows
        ]))) for i in range(len(columns))]
    total = sum(weights)
    if not isfinite(total):
        raise RenderError("Las proporciones de ancho de tabla exceden el límite numérico.")
    return [width * weight / total for weight in weights]


def _appendix_table(view: dict[str, Any], ranges: list[dict[str, Any]] | None = None) -> dict:
    rows = []
    by_id = {record["id"]: record for record in ranges or []}
    for annex in view.get("annexes", []):
        page_spec = annex.get("pages")
        selected = "Todas" if page_spec is None else ", ".join(str(p) for p in page_spec)
        record = by_id.get(annex["id"])
        rows.append([
            _text(annex["title"]), _text(annex["source_id"]), selected,
            f"{record['first_page']}–{record['last_page']}" if record else "Ver PDF ensamblado",
        ])
    return {"title": "Inventario de anexos", "columns": [
        "Anexo", "Fuente", "Páginas de origen", "Páginas del paquete"
    ], "rows": rows, "widths": [3.5, 1.5, 2, 2]}


def render_html(view: dict[str, Any], path: Path) -> None:
    """Write static UTF-8 HTML; all user-provided text is escaped."""
    status = _status(view)
    e = lambda value: escape(_text(value), quote=True)
    fragments = ["<!doctype html><html lang=\"es\"><head><meta charset=\"utf-8\">",
        '<meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'">',
        f"<title>{e(view['title'])}</title>",
        ("<style>body{font:11pt Arial,sans-serif;color:#111;max-width:1000px;margin:40px auto;padding:0 24px;line-height:1.45}"
        "h1,h2,h3{color:#000}h1{font-size:26pt}h2{font-size:17pt;margin-top:28px}h3{font-size:12pt}"
        "table{border-collapse:collapse;width:100%;table-layout:fixed;margin:10px 0 20px;font-size:10pt}"
        "th,td{border:1px solid #d9d9d9;padding:8px;text-align:left;overflow-wrap:anywhere;vertical-align:middle}"
        "th{background:#dce6f1;color:#000}tbody tr:nth-child(even){background:#f5f8fb}"
        "thead{display:table-header-group}p{white-space:pre-wrap}"
        "@media print{body{margin:0;max-width:none}tr{break-inside:avoid}h2,h3{break-after:avoid}}"
        "</style></head><body><main>"),
        f"<h1>{e(view['title'])}</h1><p>{e(view['subtitle'])}</p>",
        f"<p>Expediente {e(view['project_id'])} · Revisión {e(view['revision'])}</p>",
        f"<p><strong>{status}</strong> · Sin aprobación de ingeniería</p>"]

    def add_table(table: dict) -> None:
        columns, rows = _table_data(table)
        fragments.append(f"<h3>{e(table['title'])}</h3><table><colgroup>")
        fragments.extend(f'<col style="width:{part:.4f}%">' for part in _widths(table, 100))
        fragments.append("</colgroup><thead><tr>")
        fragments.extend(f'<th scope="col">{e(cell)}</th>' for cell in columns)
        fragments.append("</tr></thead><tbody>")
        for row in rows:
            fragments.append("<tr>" + "".join(f"<td>{e(cell)}</td>" for cell in row) + "</tr>")
        fragments.append("</tbody></table>")

    for section in view["sections"]:
        fragments.append(f"<section><h2>{e(section['title'])}</h2>")
        fragments.extend(f"<p>{e(paragraph)}</p>" for paragraph in section["paragraphs"])
        for table in section["tables"]:
            add_table(table)
        fragments.append("</section>")
    if view.get("annexes"):
        fragments.append(f"<h2>Anexos</h2><p>{_ANNEX_NOTE}</p>")
        add_table(_appendix_table(view))
    fragments.append("</main></body></html>")
    with Path(path).open("x", encoding="utf-8", newline="\n") as stream:
        stream.write("".join(fragments))


def render_docx(view: dict[str, Any], path: Path) -> None:
    """Write an editable Word body with repeating table headers and page fields."""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    status = _status(view)
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Inches(1)
    page.header_distance = page.footer_distance = Inches(0.45)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        if name == "Normal":
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(7)
            style.paragraph_format.line_spacing = 1.12
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for border in properties.findall(qn("w:pBdr")):
                properties.remove(border)
    document.styles["Title"].font.size = Pt(25)
    document.styles["Subtitle"].font.size = Pt(12)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(12)
    document.add_paragraph(_text(view["title"]), "Title")
    document.add_paragraph(_text(view["subtitle"]), "Subtitle")
    document.add_paragraph(f"Expediente {_text(view['project_id'])} · Revisión {_text(view['revision'])}")
    document.add_paragraph(f"{status} · Sin aprobación de ingeniería")
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Página ")
    for index, field_name in enumerate(("PAGE", "NUMPAGES")):
        if index:
            footer.add_run(" de ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)

    def add_table(data: dict) -> None:
        columns, rows = _table_data(data)
        document.add_paragraph(_text(data["title"]), "Heading 2")
        table = document.add_table(rows=1, cols=len(columns))
        table.autofit = False
        widths = _widths(data, 6.5)
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{side}")
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn(f"w:{key}"), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        repeated = OxmlElement("w:tblHeader")
        repeated.set(qn("w:val"), "true")
        table.rows[0]._tr.get_or_add_trPr().append(repeated)
        for row_index, values in enumerate([columns, *rows]):
            row = table.rows[0] if row_index == 0 else table.add_row()
            for column_index, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = value
                props = cell._tc.get_or_add_tcPr()
                margin = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:w"), "90")
                    edge.set(qn("w:type"), "dxa")
                    margin.append(edge)
                props.append(margin)
                if row_index == 0 or row_index % 2 == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "DCE6F1" if row_index == 0 else "F5F8FB")
                    props.append(fill)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.08
                    paragraph.paragraph_format.keep_with_next = row_index == 0
                    if row_index > 0 and len(value) <= 20:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(10 if len(columns) <= 4 else 9)
                        run.font.bold = row_index == 0
        document.add_paragraph().paragraph_format.space_after = Pt(1)

    for section in view["sections"]:
        document.add_paragraph(_text(section["title"]), "Heading 1")
        for paragraph in section["paragraphs"]:
            document.add_paragraph(_text(paragraph))
        for table in section["tables"]:
            add_table(table)
    if view.get("annexes"):
        document.add_paragraph("Anexos", "Heading 1")
        document.add_paragraph(_ANNEX_NOTE)
        add_table(_appendix_table(view))
    # python-docx otherwise overwrites silently when passed a filename.
    with Path(path).open("xb") as stream:
        document.save(stream)


def _pdf_font() -> str:
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "MemoriaSans" not in pdfmetrics.getRegisteredFontNames():
        candidates = [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path(reportlab.__file__).parent / "fonts" / "Vera.ttf",
        ]
        for candidate in candidates:
            if candidate.is_file():
                pdfmetrics.registerFont(TTFont("MemoriaSans", str(candidate)))
                break
        else:
            raise RenderError("No se encontró una fuente Unicode para el PDF.")
    return "MemoriaSans"


def _pdf_body(view: dict, ranges: list[dict], total_pages: int) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _pdf_font()
    glyphs = pdfmetrics.getFont(font).face.charToGlyph

    def safe(value: str) -> str:
        value = _text(value)
        missing = {character for character in value if not character.isspace() and ord(character) not in glyphs}
        if missing:
            raise RenderError(f"La fuente PDF no admite estos caracteres: {''.join(sorted(missing))}")
        return escape(value, quote=False).replace("\n", "<br/>")

    base = {"fontName": font, "textColor": colors.black, "leading": 14, "spaceAfter": 7}
    normal = ParagraphStyle("body", fontSize=11, **base)
    title = ParagraphStyle("title", fontSize=25, leading=30, spaceAfter=12, fontName=font)
    h1 = ParagraphStyle("h1", fontSize=16, leading=20, spaceBefore=18, spaceAfter=8, keepWithNext=True, fontName=font)
    h2 = ParagraphStyle("h2", fontSize=12, leading=15, spaceBefore=8, spaceAfter=6, fontName=font)
    cell_style = ParagraphStyle("cell", fontName=font, fontSize=9, leading=12, wordWrap="LTR", splitLongWords=True)
    story = [Paragraph(safe(view["title"]), title), Paragraph(safe(view["subtitle"]), normal),
        Paragraph(safe(f"Expediente {view['project_id']} · Revisión {view['revision']}"), normal),
        Paragraph(safe(f"{_status(view)} · Sin aprobación de ingeniería"), normal)]

    def add_table(data: dict) -> None:
        columns, rows = _table_data(data)
        caption = Paragraph(safe(data["title"]), h2)
        content = []
        for row_index, row in enumerate([columns, *rows]):
            content.append([Paragraph(safe(cell), ParagraphStyle(
                f"cell-{row_index}-{i}", parent=cell_style,
                alignment=TA_CENTER if row_index and len(cell) <= 20 else TA_LEFT,
            )) for i, cell in enumerate(row)])
        table = Table(content, colWidths=_widths(data, 468), repeatRows=1, hAlign="LEFT", splitByRow=1, splitInRow=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FB")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        # keepWithNext would keep a caption with the entire multi-page table,
        # wasting the remaining first page. Reserve caption + header + the
        # first line of the first row, then let Table split normally.
        caption_height = caption.wrap(468, 1000000)[1]
        table.wrap(468, 1000000)
        minimum = caption_height + 14 + table._rowHeights[0]
        if rows:
            minimum += min(table._rowHeights[1], 24)
        story.extend([CondPageBreak(minimum), caption, table, Spacer(1, 8)])

    for section in view["sections"]:
        story.append(Paragraph(safe(section["title"]), h1))
        story.extend(Paragraph(safe(p), normal) for p in section["paragraphs"])
        for table in section["tables"]:
            add_table(table)
    if view.get("annexes"):
        story.append(Paragraph("Anexos", h1))
        story.append(Paragraph(safe("El índice usa la paginación global del paquete. Los anexos conservan su numeración original y añaden una franja inferior con la página global."), normal))
        add_table(_appendix_table(view, ranges))

    def footer(canvas: Any, document: Any) -> None:
        if document.page > MAX_PDF_PAGES:
            raise RenderError("El cuerpo excede el límite de páginas PDF.")
        canvas.saveState()
        canvas.setFont(font, 9)
        canvas.drawRightString(540, 34, f"Página {document.page} de {total_pages}")
        canvas.restoreState()

    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=72, rightMargin=72,
        topMargin=60, bottomMargin=60, title=view["title"], author="CDCRD Digital")
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _check_annotations(annotations: Any) -> None:
    """Permit removal only when an annotation carries no rendered evidence."""
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject

    if isinstance(annotations, IndirectObject):
        annotations = annotations.get_object()
    if not isinstance(annotations, ArrayObject):
        raise RenderError("El anexo contiene anotaciones PDF no válidas.")
    for annotation in annotations:
        if isinstance(annotation, IndirectObject):
            annotation = annotation.get_object()
        if not isinstance(annotation, DictionaryObject):
            raise RenderError("El anexo contiene una anotación PDF no válida.")
        if annotation.get("/Subtype") != "/Link" or "/AP" in annotation:
            raise RenderError(
                "El anexo contiene anotaciones con apariencia o de tipo no admitido. "
                "Exporte una copia estática verificada para conservar su contenido visible."
            )
        # A link can paint a border without an appearance stream. Do not remove
        # an explicitly painted border or an unrecognized border definition.
        border = annotation.get("/Border")
        if isinstance(border, IndirectObject):
            border = border.get_object()
        if border is not None and (
            not isinstance(border, ArrayObject) or len(border) < 3 or border[2] != 0
        ):
            raise RenderError("El enlace del anexo contiene un borde visible o no reconocido.")
        border_style = annotation.get("/BS")
        if isinstance(border_style, IndirectObject):
            border_style = border_style.get_object()
        if border_style is not None and (
            not isinstance(border_style, DictionaryObject) or border_style.get("/W", 1) != 0
        ):
            raise RenderError("El enlace del anexo contiene un borde visible o no reconocido.")


def _scrub_pdf(obj: Any, seen: set[int] | None = None) -> None:
    """Remove active objects from an imported page graph, never the original file."""
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject

    seen = set() if seen is None else seen
    if isinstance(obj, IndirectObject):
        obj = obj.get_object()
    if id(obj) in seen:
        return
    seen.add(id(obj))
    if isinstance(obj, DictionaryObject):
        if "/Annots" in obj:
            _check_annotations(obj["/Annots"])
        if obj.get("/Type") == "/Filespec" or obj.get("/Subtype") == "/PS":
            raise RenderError("El anexo contiene recursos PDF externos o ejecutables.")
        if isinstance(obj, StreamObject) and "/F" in obj:
            raise RenderError("El anexo contiene un stream PDF externo.")
        for key in ("/Annots", "/AA", "/OpenAction", "/JS", "/JavaScript", "/Launch",
                    "/EmbeddedFiles", "/EF", "/AF", "/AcroForm", "/XFA", "/Collection",
                    "/RichMediaContent", "/RichMediaSettings", "/PresSteps", "/PieceInfo", "/Metadata"):
            obj.pop(key, None)
        action = obj.get("/A")
        if isinstance(action, IndirectObject):
            action = action.get_object()
        if isinstance(action, DictionaryObject) and "/S" in action:
            obj.pop("/A", None)
        for key, value in list(obj.items()):
            if key != "/Parent":
                _scrub_pdf(value, seen)
    elif isinstance(obj, ArrayObject):
        for value in obj:
            _scrub_pdf(value, seen)


def render_pdf(view: dict[str, Any], path: Path, annex_paths: dict[str, Path] | None = None) -> dict[str, Any]:
    """Assemble static pages; reject visible annotations, remove plain links."""
    from pypdf import PageObject, PdfReader, PdfWriter, Transformation
    from reportlab.pdfgen.canvas import Canvas

    _status(view)
    path = Path(path)
    if path.exists():
        raise FileExistsError(path)
    readers: dict[str, Any] = {}
    imported: list[tuple[dict, Any, list[int]]] = []
    annex_count = 0
    input_bytes = 0
    for annex in view.get("annexes", []):
        source_id = annex["source_id"]
        if annex_paths is None or source_id not in annex_paths:
            raise RenderError(f"Falta la fuente del anexo {source_id}.")
        if source_id not in readers:
            source = Path(annex_paths[source_id])
            if source.stat().st_size > MAX_PDF_BYTES:
                raise RenderError("El anexo excede el límite de tamaño PDF.")
            with source.open("rb") as stream:
                data = stream.read(MAX_PDF_BYTES + 1)
            if len(data) > MAX_PDF_BYTES:
                raise RenderError("El anexo excede el límite de tamaño PDF.")
            input_bytes += len(data)
            if input_bytes > MAX_TOTAL_PDF_BYTES:
                raise RenderError("Los anexos exceden el límite de tamaño total PDF.")
            try:
                reader = PdfReader(BytesIO(data), strict=True)
                if reader.is_encrypted:
                    raise RenderError("Los anexos cifrados no están admitidos.")
                if len(reader.pages) > MAX_PDF_PAGES:
                    raise RenderError("El anexo excede el límite de páginas PDF.")
            except RenderError:
                raise
            except Exception as exc:
                raise RenderError(f"No se puede leer el anexo PDF {source_id}.") from exc
            readers[source_id] = reader
        reader = readers[source_id]
        requested = annex.get("pages")
        pages = list(range(1, len(reader.pages) + 1)) if requested is None else requested
        if not pages or any(type(n) is not int or n < 1 or n > len(reader.pages) for n in pages):
            raise RenderError(f"Selección de páginas no válida para {source_id}.")
        if len(set(pages)) != len(pages):
            raise RenderError("Un anexo no puede repetir páginas de origen.")
        annex_count += len(pages)
        if annex_count >= MAX_PDF_PAGES:
            raise RenderError("El paquete excede el límite de páginas PDF.")
        imported.append((annex, reader, pages))

    body_pages = 1
    for _ in range(6):
        ranges = []
        next_page = body_pages + 1
        for annex, _, pages in imported:
            ranges.append({"id": annex["id"], "source_id": annex["source_id"], "title": annex["title"],
                           "first_page": next_page, "last_page": next_page + len(pages) - 1})
            next_page += len(pages)
        total_pages = body_pages + annex_count
        if total_pages > MAX_PDF_PAGES:
            raise RenderError("El paquete excede el límite de páginas PDF.")
        body_data = _pdf_body(view, ranges, total_pages)
        body = PdfReader(BytesIO(body_data))
        actual_pages = len(body.pages)
        if actual_pages == body_pages:
            break
        body_pages = actual_pages
    else:
        raise RenderError("No convergió la paginación del índice de anexos.")

    writer = PdfWriter()
    for page in body.pages:
        writer.add_page(page)
    writer.add_outline_item("Cuerpo de la memoria", 0)
    for (annex, reader, selected), record in zip(imported, ranges):
        for number in selected:
            original = reader.pages[number - 1]
            # Only mutate the in-memory reader. Source files are never opened for write.
            _scrub_pdf(original)
            if original.rotation:
                rotation_writer = PdfWriter()
                original = rotation_writer.add_page(original)
                original.transfer_rotation_to_content()
            box = original.cropbox
            width, height = float(box.width), float(box.height)
            if not all(isfinite(v) and 72 <= v <= 14400 for v in (width, height)):
                raise RenderError("El tamaño de página del anexo no está admitido.")
            band = 24.0
            combined = PageObject.create_blank_page(width=width, height=height + band)
            combined.merge_transformed_page(original, Transformation().translate(-float(box.left), band - float(box.bottom)))
            overlay_data = BytesIO()
            overlay = Canvas(overlay_data, pagesize=(width, height + band))
            overlay.setFont(_pdf_font(), 8)
            overlay.drawCentredString(width / 2, 8, f"Página {len(writer.pages) + 1} de {total_pages}")
            overlay.save()
            combined.merge_page(PdfReader(BytesIO(overlay_data.getvalue())).pages[0])
            writer.add_page(combined, excluded_keys=["/Annots", "/AA", "/Metadata"])
        # Destination pages must already belong to the writer; an index for a
        # future page is serialized as a number instead of a page reference.
        writer.add_outline_item(annex["title"], record["first_page"] - 1)
    writer.add_metadata({"/Title": _text(view["title"]), "/Author": "CDCRD Digital", "/Subject": "Memoria pendiente de aprobación de ingeniería"})
    with path.open("xb") as stream:
        writer.write(stream)
    return {"body_pages": body_pages, "total_pages": total_pages, "annexes": ranges}
