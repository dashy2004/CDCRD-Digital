"""Un mismo resultado PM llega al visor y al Excel, sin sustituirlo por FEM."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from plugin_losas.exportacion_pm import escribir_exportacion_pm
from plugin_losas.resultados_pm import dto_visor


class TestExportacionPM(unittest.TestCase):
    def test_momentos_y_acero_del_visor_son_los_exportados(self):
        proyecto = {"losas": [{"mark": "EJ1", "contorno": [[0, 0], [4, 0], [4, 5], [0, 5]]}]}
        resultados = {"panos": [{"mark": "EJ1", "nivel": "=2+2",
            "cargas_tm2": {"q_ultima": 1.23456}, "espesor": {"h_usado_m": 0.2},
            "franja_x": {"as_diseno": 9999, "disponer": "FEM diferente"}}],
            "pieper_martens": {"losas": {"EJ1": {"codigo_dl": 10,
                "mfx": 1.1111111, "mfy": 2.2222222, "msx": 3.3333333, "msy": 4.4444444,
                "q_pm_tm2": 1.235}}, "apoyos": []}}
        armados = {k: {"as_requerido": 100 + i, "as_minimo": 200, "as_diseno": 200 + i,
                      "numero_barra": 3, "espaciamiento": 0.25, "as_provista": 284,
                      "disponer": "Barra de ejemplo", "cumple": True, "seccion_insuficiente": False}
                   for i, k in enumerate(("mfx", "mfy", "msx", "msy"))}
        acero = {"losas": {"EJ1": {"armados": armados}}, "apoyos": []}
        dto = dto_visor(proyecto, resultados, acero=acero)
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "pm.xlsx")
            escribir_exportacion_pm(proyecto, resultados, path, acero=acero)
            wb = load_workbook(path, data_only=True)
            row = list(wb["Momentos PM"].values)[1]
            self.assertEqual(wb["Momentos PM"]["B2"].data_type, "s")
            self.assertEqual(row[4], 1.235)
            for i, key in enumerate(("mfx", "mfy", "msx", "msy"), 6):
                self.assertEqual(row[i], dto.campos[key]["valores"]["EJ1"])
            self.assertEqual(wb["Acero PM por metro"]["G2"].value, dto.campos["as_x"]["valores"]["EJ1"])
            self.assertEqual(dto.campos["as_apoyo"]["valores"]["EJ1"], 203)
            self.assertNotEqual(dto.campos["as_x"]["valores"]["EJ1"], 9999)
            wb.close()

    def test_acero_pm_ausente_se_muestra_pendiente(self):
        resultados = {"panos": [{"mark": "EJ1", "cargas_tm2": {"q_ultima": 1.23456}}],
                      "pieper_martens": {"losas": {"EJ1": {
                          "codigo_dl": 10, "mfx": 0, "mfy": 0, "msx": 0, "msy": 0}}}}
        dto = dto_visor({"losas": []}, resultados)
        self.assertIsNone(dto.campos["as_x"]["valores"]["EJ1"])
        self.assertEqual(dto.panos[0]["armaduras"]["x"], "PM pendiente")
        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "pm.xlsx")
            escribir_exportacion_pm({"losas": []}, resultados, path)
            wb = load_workbook(path, data_only=True)
            self.assertIsNone(wb["Momentos PM"]["E2"].value)
            self.assertIsNone(wb["Acero PM por metro"]["G2"].value)
            wb.close()
