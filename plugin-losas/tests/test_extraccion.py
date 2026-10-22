"""Filas del panel: extracción desde el export de Revit, uso editable y exports."""
import os
import tempfile
import unittest

from puente_losas.contrato import aplicar_usos, panos_desde_dict

from plugin_losas.extraccion import (asignar_uso, escribir_tabla_excel,
                                     escribir_usos_json, filas_desde_export,
                                     tabla_filas, usos_para_puente)
from plugin_losas.usos_tabla4 import cargar_tabla4

# Export inventado con la forma del bloque "losas" de exportar_proyecto
# (revit-mcp-tools). Ningún dato corresponde a un proyecto real.
EXPORT_INVENTADO = {
    "meta": {"proyecto": "Edificio Ejemplo PL", "unidades": {"longitud": "m"}},
    "losas": [
        {"id": "L1", "mark": "LE1", "nivel": "N1 +3.00", "z_tope": 3.0,
         "contorno": [[0, 0], [5, 0], [5, 6], [0, 6]], "huecos": [],
         "lx": 5.0, "ly": 6.0, "area_m2": 30.0,
         "espesor": {"h_usado": 0.25}, "tipo_revit": "Generic - 0.25", "revit_id": 1001},
        {"id": "L2", "mark": "LE2", "nivel": "N1 +3.00", "z_tope": 3.0,
         "contorno": [[5, 0], [10, 0], [10, 6], [5, 6]], "huecos": [],
         "lx": 5.0, "ly": 6.0, "area_m2": 30.0,
         "espesor": {"h_usado": 0.25}, "tipo_revit": "Generic - 0.25", "revit_id": 1002},
        {"id": "L3", "mark": "LE6", "nivel": "N1 +3.00", "z_tope": 3.0,
         "contorno": [[-1.8, 0], [0, 0], [0, 6], [-1.8, 6]], "huecos": [],
         "lx": 1.8, "ly": 6.0, "area_m2": 10.8,
         "espesor": {"h_usado": 0.20}, "tipo_revit": "Generic - 0.20", "revit_id": 1003},
    ],
}

ID_AULAS = "escuelas-y-colegios-aulas"
ID_BALCON = "balcones-y-terrazas-balcones-y-terrazas"
ID_ESCENARIO = "edificios-para-reuniones-escenarios"


class TestFilas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tabla = cargar_tabla4()

    def setUp(self):
        self.filas = filas_desde_export(EXPORT_INVENTADO)

    def test_extraccion_numeracion_y_dimensiones(self):
        self.assertEqual([f.mark for f in self.filas], ["LE1", "LE2", "LE6"])
        f = self.filas[0]
        self.assertEqual((f.lx, f.ly, f.area_m2, f.espesor_m), (5.0, 6.0, 30.0, 0.25))
        self.assertEqual(f.nivel, "N1 +3.00")
        self.assertEqual(f.fuente, "revit")
        self.assertEqual(f.revit_id, 1001)
        self.assertIsNone(f.uso_tabla4)

    def test_asignar_uso_valido_es_inmutable(self):
        nuevas = asignar_uso(self.filas, "LE1", ID_AULAS, self.tabla)
        self.assertEqual(nuevas[0].uso_tabla4, ID_AULAS)
        self.assertIsNone(self.filas[0].uso_tabla4)      # la lista original no cambia

    def test_asignar_uso_inexistente_falla(self):
        with self.assertRaises(KeyError):
            asignar_uso(self.filas, "LE1", "uso-que-no-existe", self.tabla)
        with self.assertRaises(KeyError):
            asignar_uso(self.filas, "L999", ID_AULAS, self.tabla)

    def test_usos_para_puente_y_avisos(self):
        filas = asignar_uso(self.filas, "LE1", ID_AULAS, self.tabla)
        filas = asignar_uso(filas, "LE2", ID_ESCENARIO, self.tabla)
        filas = asignar_uso(filas, "LE6", ID_BALCON, self.tabla)
        usos, avisos = usos_para_puente(filas, self.tabla)

        self.assertEqual(usos["LE1"], {"uso": "entrepisos"})
        self.assertEqual(usos["LE2"]["uso"], "gimnasio")     # 7.18 kN: el más pesado
        self.assertTrue(any("LE2" in a and "supera el set" in a for a in avisos))
        self.assertEqual(usos["LE6"], {"es_vuelo": True})    # balcón: regla, sin set
        self.assertTrue(any("LE6" in a and "es_vuelo" in a for a in avisos))

    def test_fila_sin_uso_genera_aviso_y_no_entra(self):
        usos, avisos = usos_para_puente(self.filas, self.tabla)
        self.assertEqual(usos, {})
        self.assertEqual(len(avisos), 3)

    def test_usos_json_lo_acepta_el_puente(self):
        """Integración: el usos.json generado pasa por aplicar_usos de puente_losas."""
        filas = asignar_uso(self.filas, "LE1", ID_AULAS, self.tabla)
        filas = asignar_uso(filas, "LE6", ID_BALCON, self.tabla)
        usos, _ = usos_para_puente(filas, self.tabla)
        panos = aplicar_usos(panos_desde_dict(EXPORT_INVENTADO), usos)
        por_mark = {p.mark: p for p in panos}
        self.assertEqual(por_mark["LE1"].uso, "entrepisos")
        self.assertTrue(por_mark["LE6"].es_vuelo)

    def test_tabla_y_excel(self):
        filas = asignar_uso(self.filas, "LE1", ID_AULAS, self.tabla)
        contenido = tabla_filas(filas, self.tabla)
        self.assertEqual(len(contenido), 3)
        self.assertEqual(contenido[0][0], "LE1")
        self.assertIn("Aulas", contenido[0])
        self.assertEqual(contenido[1][7], "(sin asignar)")

        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, "tabla_losas.xlsx")
            escribir_tabla_excel(filas, self.tabla, ruta, proyecto="Ejemplo Inventado")
            ws = load_workbook(ruta)["Losas Uso"]
            self.assertIn("Ejemplo Inventado", ws["A1"].value)
            self.assertEqual(ws.cell(row=2, column=1).value, "Mark")
            self.assertEqual(ws.cell(row=4, column=1).value, "LE1")
            self.assertEqual(ws.max_row, 6)

            ruta_usos = os.path.join(d, "usos.json")
            usos, _ = usos_para_puente(filas, self.tabla)
            escribir_usos_json(usos, ruta_usos)
            self.assertTrue(os.path.getsize(ruta_usos) > 0)


if __name__ == "__main__":
    unittest.main()
