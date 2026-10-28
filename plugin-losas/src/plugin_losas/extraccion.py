"""Núcleo del panel de losas: filas (numeración + dimensiones + uso) y exports.

La misma lógica sirve a las dos puntas:

- **Revit**: las filas salen del export de ``exportar_proyecto``
  (``revit-mcp-tools``, marker "MCP:LOSA-ESQ"), vía el contrato ``PanoLosa``
  de ``puente_losas`` — no se re-parsea nada.
- **AutoCAD**: las filas salen de ``acad_dxf.losas_desde_dxf`` (recuadros +
  numeración del DWG exportado a DXF).

El uso editable es un id de la Tabla 4 (``usos_tabla4``). De las filas salen:
la tabla Excel de losas, y el ``usos.json`` que consume el CLI de
``puente-losas`` (que a su vez reusa ``informe_excel.py`` — el cálculo y el
informe completos no se reimplementan aquí).
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field, replace

from puente_losas.contrato import panos_desde_dict

from plugin_losas.usos_tabla4 import SugerenciaSet, TablaUsos, set_sugerido

FUENTES_VALIDAS = ("revit", "acad")
GRUPO_BALCONES = "Balcones y terrazas"


@dataclass(frozen=True)
class FilaLosa:
    """Una fila del panel: lo que el usuario ve y edita por losa."""
    mark: str
    nivel: str | None
    lx: float
    ly: float
    area_m2: float | None
    espesor_m: float | None
    fuente: str                      # "revit" | "acad"
    uso_tabla4: str | None = None    # id de usos_tabla4 (editable en el panel)
    es_vuelo: bool = False
    revit_id: int | None = None
    contorno: tuple[tuple[float, float], ...] = field(default_factory=tuple)
    criterio_espesor: str | None = None
    espesor_candidato_m: float | None = None
    espesor_referencia_m: float | None = None

    def __post_init__(self) -> None:
        if self.fuente not in FUENTES_VALIDAS:
            raise ValueError(f"{self.mark}: fuente '{self.fuente}' no está en {FUENTES_VALIDAS}.")
        if self.lx <= 0 or self.ly <= 0:
            raise ValueError(f"{self.mark}: dimensiones no positivas (lx={self.lx}, ly={self.ly}).")


def filas_desde_export(d: dict) -> list[FilaLosa]:
    """Filas del panel desde el dict de ``exportar_proyecto`` (bloque ``losas``)."""
    filas = []
    for p in panos_desde_dict(d):
        filas.append(FilaLosa(
            mark=p.mark, nivel=p.nivel, lx=p.lx, ly=p.ly,
            area_m2=p.area_m2, espesor_m=p.espesor or None,
            fuente="revit", es_vuelo=p.es_vuelo, revit_id=p.revit_id,
            contorno=tuple((x, y) for x, y in p.contorno)))
    return filas


def asignar_uso(filas: list[FilaLosa], mark: str, uso_id: str | None,
                tabla: TablaUsos) -> list[FilaLosa]:
    """Devuelve una lista nueva con el uso del Mark cambiado (None lo limpia)."""
    if uso_id is not None:
        tabla.por_id(uso_id)                      # valida: KeyError si no existe
    if mark not in {f.mark for f in filas}:
        raise KeyError(f"mark '{mark}' no está en las filas del panel.")
    return [replace(f, uso_tabla4=uso_id) if f.mark == mark else f for f in filas]


@dataclass(frozen=True)
class FilaResuelta:
    """Fila + uso de la Tabla 4 resuelto a set de carga del flujo Excel/ETABS."""
    fila: FilaLosa
    sugerencia: SugerenciaSet | None   # None = sin uso asignado


def resolver_filas(filas: list[FilaLosa], tabla: TablaUsos) -> list[FilaResuelta]:
    out = []
    for f in filas:
        if f.uso_tabla4 is None:
            out.append(FilaResuelta(f, None))
        else:
            out.append(FilaResuelta(f, set_sugerido(tabla.por_id(f.uso_tabla4))))
    return out


def usos_para_puente(filas: list[FilaLosa], tabla: TablaUsos
                     ) -> tuple[dict[str, dict], list[str]]:
    """El ``usos.json`` que consume ``puente-losas --usos`` + avisos de la resolución.

    Por fila con uso asignado: ``{"uso": <set>}`` (y ``es_vuelo`` si aplica).
    Avisos cuando el set no cubre la viva de la Tabla 4, cuando el uso es una
    regla (balcones) o cuando una fila quedó sin uso: nada se resuelve en
    silencio.
    """
    usos: dict[str, dict] = {}
    avisos: list[str] = []
    for r in resolver_filas(filas, tabla):
        f = r.fila
        if r.sugerencia is None:
            avisos.append(f"{f.mark}: sin uso asignado; conserva el uso por defecto del puente.")
            continue
        uso = tabla.por_id(f.uso_tabla4)
        entrada: dict = {}
        if r.sugerencia.set_carga is None:
            avisos.append(f"{f.mark}: '{uso.uso}' no tiene set directo — {r.sugerencia.nota}.")
        else:
            entrada["uso"] = r.sugerencia.set_carga
            if not r.sugerencia.cubre:
                avisos.append(f"{f.mark}: {r.sugerencia.nota}.")
        if f.es_vuelo or uso.grupo == GRUPO_BALCONES:
            entrada["es_vuelo"] = True
            if uso.grupo == GRUPO_BALCONES and not f.es_vuelo:
                avisos.append(f"{f.mark}: uso de balcón/terraza — se marca es_vuelo "
                              "(verificar que el paño realmente vuela).")
        if entrada:
            usos[f.mark] = entrada
    return usos, avisos


def escribir_usos_json(usos: dict[str, dict], ruta: str) -> None:
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(usos, f, indent=1, ensure_ascii=False)


# ------------------------- tabla Excel de losas -------------------------
ENCABEZADOS_TABLA = [
    ("Mark", ""), ("Nivel", ""), ("Fuente", ""),
    ("Lx", "m"), ("Ly", "m"), ("Área", "m²"), ("Espesor", "m"),
    ("Uso Tabla 4", ""), ("Grupo", ""), ("Viva Tabla 4", "kN/m²"),
    ("Viva Tabla 4", "t/m²"), ("Set de carga", ""), ("Viva del set", "t/m²"),
    ("Set cubre", ""), ("Vuelo", ""), ("Nota", ""),
]


def tabla_filas(filas: list[FilaLosa], tabla: TablaUsos) -> list[list]:
    """Las filas de la tabla Excel (valores planos, en el orden de ENCABEZADOS_TABLA)."""
    out = []
    for r in resolver_filas(filas, tabla):
        f = r.fila
        if r.sugerencia is None:
            uso_cols = ["(sin asignar)", "", "", "", "", "", "", ""]
        else:
            u = tabla.por_id(f.uso_tabla4)
            s = r.sugerencia
            uso_cols = [
                u.uso, u.grupo,
                round(u.viva_kn_m2, 2) if u.viva_kn_m2 is not None else "regla",
                round(u.viva_tm2, 3) if u.viva_tm2 is not None else "",
                s.set_carga or "(declarar)",
                round(s.viva_set_tm2, 2) if s.viva_set_tm2 is not None else "",
                "SÍ" if s.cubre else "NO",
                s.nota or (u.regla or ""),
            ]
        vuelo = "SÍ" if f.es_vuelo else "NO"
        out.append([f.mark, f.nivel or "", f.fuente,
                    round(f.lx, 3), round(f.ly, 3),
                    round(f.area_m2, 2) if f.area_m2 is not None else "",
                    round(f.espesor_m, 3) if f.espesor_m is not None else "",
                    *uso_cols[:6], uso_cols[6], vuelo, uso_cols[7]])
    return out


def escribir_tabla_excel(filas: list[FilaLosa], tabla: TablaUsos, ruta: str,
                         proyecto: str = "Proyecto Ejemplo") -> None:
    """Tabla de losas (numeración + dimensiones + uso) para el Excel de losas.

    Es la tabla de entrada del flujo; el informe de cálculo completo lo escribe
    ``puente_losas.informe_excel`` al correr el CLI con el ``usos.json`` de
    ``usos_para_puente`` (reuso, no reimplementación).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Losas Uso"
    ws.append([f"Tabla de losas — {proyecto} — plugin-losas "
               "(usos Tabla 4 CDCRD; sets del flujo de la Ing. Arlette Ramírez)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([h for h, _ in ENCABEZADOS_TABLA])
    ws.append([u for _, u in ENCABEZADOS_TABLA])
    negrita = Font(bold=True)
    centro = Alignment(horizontal="center")
    for celda in ws[2] + ws[3]:
        celda.font = negrita
        celda.alignment = centro
    rojo = PatternFill("solid", start_color="FFC7CE")
    col_cubre = next(i for i, (h, _) in enumerate(ENCABEZADOS_TABLA, start=1) if h == "Set cubre")
    for fila in tabla_filas(filas, tabla):
        ws.append(fila)
        celda = ws.cell(row=ws.max_row, column=col_cubre)
        if celda.value == "NO":
            celda.fill = rojo
    ws.freeze_panes = "A4"
    wb.save(ruta)
