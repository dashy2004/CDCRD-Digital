"""Exporta la misma instantánea Pieper–Martens que consume el visor.

Los resultados FEM no se usan como sustitutos de resultados PM ausentes.
El acero comercial corresponde a una franja maciza de 1 m; no diseña nervios.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font

from plugin_losas.resultados_pm import dto_visor


def escribir_exportacion_pm(proyecto: dict, resultados: dict, ruta: str,
                           nombre: str = "Proyecto Ejemplo",
                           acero: dict | None = None) -> None:
    """Vuelca valores PM sin redondear ni recalcularlos en el exportador."""
    dto = dto_visor(proyecto, resultados, nombre, acero=acero)
    panos = {p["mark"]: p for p in resultados.get("panos", [])}
    pm = resultados["pieper_martens"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Momentos PM"
    ws.append(["Mark", "Nivel", "Código PM", "q última tf/m²", "q PM tf/m²",
               "h usado m", "Mfx tf·m/m", "Mfy tf·m/m", "Msx tf·m/m",
               "Msy tf·m/m", "Estado / hipótesis"])
    for p in dto.panos:
        mark = p["mark"]
        r = panos.get(mark, {})
        q = (r.get("cargas_tm2") or {}).get("q_ultima")
        ws.append([mark, r.get("nivel"), p["codigo_dl"], q,
                   pm["losas"][mark].get("q_pm_tm2"),
                   dto.campos["h_usado"]["valores"][mark],
                   *[dto.campos[k]["valores"][mark] for k in ("mfx", "mfy", "msx", "msy")],
                   p["nota"] or "Apoyos inferidos; revisión pendiente"])
    barras = wb.create_sheet("Acero PM por metro")
    barras.append(["Mark", "Zona PM", "Mu tf·m/m", "Base de sección", "As requerido mm²/m",
                   "As mínimo mm²/m", "As diseño mm²/m", "Barra #", "Separación m",
                   "As provista mm²/m", "Propuesta de franja", "Cumple franja",
                   "Sección insuficiente", "Estado nervios"])
    for mark, momentos in pm["losas"].items():
        armados = ((acero or {}).get("losas", {}).get(mark) or {}).get("armados", {})
        for zona in ("mfx", "mfy", "msx", "msy"):
            a = armados.get(zona, {})
            barras.append([mark, zona, momentos.get(zona), "Franja rectangular maciza de 1 m",
                           a.get("as_requerido"), a.get("as_minimo"), a.get("as_diseno"),
                           a.get("numero_barra"), a.get("espaciamiento"), a.get("as_provista"),
                           a.get("disponer"), a.get("cumple"), a.get("seccion_insuficiente"),
                           "No es armado por nervio; requiere comprobación de sección"])
    apoyos = wb.create_sheet("Apoyos PM")
    apoyos.append(["Losa I", "Losa J", "Dirección", "Mu I tf·m/m", "Mu J tf·m/m",
                   "Mu balanceado tf·m/m", "Modo"])
    for a in pm.get("apoyos", []):
        apoyos.append([a.get(k) for k in ("losa_i", "losa_j", "direccion", "mu_i", "mu_j", "mu_ij", "modo")])
    notas = wb.create_sheet("Alcance")
    for row in [["Proyecto", nombre], ["Fuente", "resultados.json y acero_pm.json de la misma corrida"],
                ["Momentos", "Pieper–Martens; valores crudos sin sustituir por FEM"],
                ["Carga PM", "El motor redondea q última a 0.001 tf/m² para el archivo DL"],
                ["Acero", "Franja maciza de 1 m, valores del cálculo PM; nervios pendientes"],
                ["Aprobación", "Resultados preliminares; no constituye diseño aprobado"],
                *[["Crédito", c] for c in dto.creditos]]:
        notas.append(row)
    for hoja in wb:
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for row in hoja:
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                if isinstance(cell.value, str):
                    # Frozen exports contain data, never executable formulas.
                    cell.data_type = "s"
        for cell in hoja[1]:
            cell.font = Font(name="Arial", bold=True)
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = min(55, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(ruta)
